# -*- coding: utf-8 -*-
"""Word do relatorio do Painel da Diretoria (v0.63.0): mesma estrutura do PDF, texto editavel."""
import io

from docx import Document

import relatorio_diretoria_docx as RDX
import relatorio_diretoria_dados as RDD


def _dados():
    els = [RDD.kpis([RDD.K("Receita", "R$ 1,0 MM", "no mês"), RDD.K("Custo", "−R$ 0,5 MM")], 2),
           RDD.tab(["Conta", "Valor"], [("norm", ["Seguro", "1.234,50"]), ("tot", ["Total", ("−2.000,00", None)])], [3, 1], titulo="Despesas"),
           RDD.barras(["jan", "fev", "mar"], [1e6, 2.5e6, None], "mi", "R$ MM", "Evolução"),
           RDD.sec("Subtítulo"), RDD.nota("Filtros do painel: conta Sette.")]
    return {"meta": {"per_rot": "Jan–Set/2026", "gerado_em": "30/10/2026 15:00", "usuario": "Rafael", "filtros": ["Período: Jan–Set/2026", "Base: Caixa"], "mes": "2026-09", "n_blocos": 1},
            "abas": [{"aba": "desp", "rotulo": "Despesas", "blocos": [{"key": "x", "titulo": "Despesas por conta", "els": els}], "obs": "Conferir com a Suellen."}]}


def test_word_tem_texto_editavel_tabelas_e_observacao():
    b = RDX.gerar_docx(_dados())
    assert b[:2] == b"PK"
    d = Document(io.BytesIO(b))
    txt = "\n".join(p.text for p in d.paragraphs)
    assert "Painel da Diretoria" in txt and "setembro de 2026" in txt and "Base: Caixa" in txt
    assert any("JAN–SET/2026" in c for t in d.tables for r in t.rows for c in [x.text for x in r.cells])           # etiqueta do periodo na capa
    assert "Despesas por conta" in txt and "Subtítulo" in txt
    celulas = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "Filtros do painel: conta Sette." in celulas and "1.234,50" in celulas and "−2.000,00" in celulas and any(c.startswith("RECEITA") and "R$ 1,0 MM" in c for c in celulas)
    assert "1,0" in celulas and "2,5" in celulas and "–" in celulas          # barras viram tabela de valores
    assert any("Conferir com a Suellen." in c for c in celulas)


def test_word_sem_abas_nao_quebra():
    d = Document(io.BytesIO(RDX.gerar_docx({"meta": {"per_rot": "x"}, "abas": []})))
    assert "Painel da Diretoria" in "\n".join(p.text for p in d.paragraphs)
