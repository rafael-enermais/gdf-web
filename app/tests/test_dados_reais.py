# -*- coding: utf-8 -*-
"""
Testes com DADOS REAIS da Anastacio (agosto/2026) -- os arquivos NAO ficam no repositorio (publico).
Rodar:  GDF_DADOS_REAIS=<pasta> pytest -q app/tests/test_dados_reais.py
A pasta precisa ter: Anastacio_Demonstrativos_Bancos_08.2026_1.xlsx, gdf_balancetes_2026_lidos_v0.1.json (balancetes mensais
lidos dos PDFs) e Relatorios_Contabeis_Balancete_(Texto)_Balancete_-_Debito_Credito.csv (acumulado 01-08).
Prova: o motor reproduz Balanco (dez/25, jul, ago), DRE e os 17 indicadores do Excel a centavo; o CSV acumulado concilia.
"""
import json
import os
from pathlib import Path

import openpyxl
import pytest

import importador_csv as I
import motor as M

pytestmark = pytest.mark.dados_reais
TOL = 0.01
NOME_XLSX = "Anastacio_Demonstrativos_Bancos_08.2026_1.xlsx"
NOME_LIDOS = "gdf_balancetes_2026_lidos_v0.1.json"
NOME_CSV = "Relatorios_Contabeis_Balancete_(Texto)_Balancete_-_Debito_Credito.csv"


@pytest.fixture(scope="module")
def pasta():
    p = os.environ.get("GDF_DADOS_REAIS")
    if not p or not Path(p).exists():
        pytest.skip("defina GDF_DADOS_REAIS=<pasta com os dados reais>")
    return Path(p)


@pytest.fixture(scope="module")
def b(pasta):
    d = json.load(open(pasta / NOME_LIDOS))["periodos"]
    return M.Balancetes({k: v["contas"] for k, v in d.items() if k.startswith("2026-")})


@pytest.fixture(scope="module")
def lidos(pasta):
    return json.load(open(pasta / NOME_LIDOS))["periodos"]


@pytest.fixture(scope="module")
def wb(pasta):
    return openpyxl.load_workbook(pasta / NOME_XLSX, data_only=True)


@pytest.fixture(scope="module")
def csvd(pasta):
    return I.ler_csv(str(pasta / NOME_CSV))


@pytest.fixture(scope="module")
def mensal(b, lidos):
    return b, lidos

# linha do Excel -> linha do motor
BP_ROWS = {11: "dep_vista", 12: "aplic", 13: "caixa_eq", 14: "adiant", 15: "trib_rec", 16: "desp_ant_cp", 17: "ac",
           20: "dep_jud", 21: "desp_ant_lp", 22: "rlp", 23: "conc_bruto", 24: "conc_red", 25: "conc_liq", 26: "anc", 28: "ativo",
           33: "bndes_cp", 34: "fornec", 35: "prov_constr", 36: "imp_rec", 37: "trib_ret", 38: "pc",
           41: "bndes_lp", 42: "trib_dif", 43: "pnc", 46: "capital", 47: "afac", 48: "res_legal", 49: "res_ret",
           50: "prej", 51: "pl", 53: "passivo_pl"}
DRE_ROWS = {10: "pis", 11: "cofins", 12: "deducoes", 13: "rec_liq", 15: "custo_constr", 16: "fretes", 17: "custos_serv", 18: "res_bruto",
            21: "serv_prof", 22: "cartorio", 23: "seguro", 24: "importacao", 25: "ajustes", 26: "desp_adm", 27: "ebit",
            30: "rend_aplic", 31: "rec_nt", 32: "descontos", 33: "rec_fin", 34: "desp_banc", 35: "juros", 36: "desp_fin",
            37: "res_fin", 39: "res_antes_ir", 40: "ir_cs", 42: "res_liq"}
IND_ROWS = {11: "liq_corrente", 12: "liq_imediata", 13: "liq_geral", 14: "ccl", 17: "div_bruta", 18: "caixa_neg", 19: "div_liquida",
            20: "div_cp_sobre_bruta", 21: "div_bruta_ativo", 22: "div_liq_sobre_concessao", 25: "endiv_geral", 26: "comp_endiv_cp",
            27: "pl_ativo", 28: "capital_aportado", 31: "res_fin", 32: "ebit_ebitda", 33: "res_liq"}

def _num(v):
    return None if v in (None, "-") else float(v)

def test_balanco(b, wb):
    ws = wb["Balanço Patrimonial"]; bp = M.balanco(b, "2026-08", "2026-07")
    erros = []
    for col, nome in (("C", "abertura"), ("D", "mes_ant"), ("E", "mes_ref")):
        for r, k in BP_ROWS.items():
            esp = _num(ws[f"{col}{r}"].value); got = bp[nome][k]
            if esp is None or abs(esp - got) > TOL:
                erros.append((nome, k, esp, got))
    assert not erros, erros

def test_dre(b, wb):
    ws = wb["DRE"]; d = M.dre(b, "2026-08"); erros = []
    for col, nome in (("C", "ate_mes_ant"), ("D", "mes"), ("E", "acumulado")):
        for r, k in DRE_ROWS.items():
            esp = _num(ws[f"{col}{r}"].value); got = d[nome][k]
            if esp is None or abs(esp - got) > TOL:
                erros.append((nome, k, esp, got))
    assert not erros, erros

def test_indicadores(b, wb):
    ws = wb["Indicadores"]; bp = M.balanco(b, "2026-08", "2026-07"); d = M.dre(b, "2026-08"); erros = []
    for col, nome, dre_ in (("C", "abertura", None), ("D", "mes_ant", d["ate_mes_ant"]), ("E", "mes_ref", d["acumulado"])):
        ind = M.indicadores(bp[nome], dre_)
        for r, k in IND_ROWS.items():
            esp = _num(ws[f"{col}{r}"].value); got = ind[k]
            if esp is None and got is None:
                continue
            if esp is None or got is None or abs(esp - got) > (0.01 if abs(esp) > 100 else 1e-9):
                erros.append((nome, k, esp, got))
    assert not erros, erros

def test_conferencias(b):
    falhas = [c for c in M.conferencias(b) if not c["ok"]]
    assert not falhas, falhas

def test_ajuste_agosto(b):
    assert abs(b.ajuste_mes("2026-08") - 7.40) < 0.005
    assert all(b.ajuste_mes(m) == 0.0 for m in b.ordem[:-1])

def test_soma_mensal_igual_balancete_acumulado(lidos):
    """Soma dos 7 balancetes mensais = balancete acumulado 01 a 07 (por chave de resultado)."""
    d = lidos
    mens = M.Balancetes({k: v["contas"] for k, v in d.items() if k.startswith("2026-") and k <= "2026-07"})
    acum = M.Balancetes({"ACUM": d["ACUM-01a07"]["contas"]})
    for k in M.CHAVES_DRE:
        soma = round(sum(mens.mov(m, k) for m in mens.ordem), 2)
        assert abs(soma - acum.mov("ACUM", k)) <= 0.011, (k, soma, acum.mov("ACUM", k))

def test_conferencias_quantidade(b):
    r = M.conferencias(b)
    assert len(r) == 8 * 8 - 1          # 8 meses x 8 conferencias; "Continuidade" nao existe no 1o mes


def test_cabecalho(csvd):
    cab, contas = csvd
    assert cab["cnpj"] == "54.800.488/0001-60" and (cab["tipo"], cab["periodo"]) == ("ACUMULADO", "2026-08") and str(cab["ini"]) == "2026-01-01" and str(cab["fim"]) == "2026-08-31"
    assert len(contas) == 251 and not cab["ignoradas"]

def test_linhas_fecham(csvd):
    _, c = csvd
    ruins = [x["cl"] for x in c if x["cl"] != "8" and not (abs(x["ant"] + x["deb"] - x["cred"] - x["sal"]) < 0.005 or abs(x["ant"] - x["deb"] + x["cred"] - x["sal"]) < 0.005)]
    assert not ruins, ruins

def test_resultado_acumulado_igual_motor(csvd, mensal):
    """Soma dos 8 balancetes mensais (motor) = balancete acumulado 01-08. Unica diferenca: despesas bancarias = R$ 7,40 (= 'Outros ajustes liquidos')."""
    _, c = csvd; b, _ = mensal
    for k in M.CHAVES_DRE:
        motor = round(sum(b.mov(m, k) for m in b.ordem), 2)
        csv_ = round(sum(x["cred"] - x["deb"] for x in c if M.chave_da_conta(x["cl"]) == k and not x["sint"]), 2)
        if k == "desp_banc":
            assert abs((csv_ - motor) - b.ajuste_mes("2026-08")) < 0.011, (motor, csv_)
        else:
            assert abs(motor - csv_) < 0.011, (k, motor, csv_)

def test_linha8_acumulada_igual_dre(csvd, mensal):
    _, c = csvd; b, _ = mensal
    l8 = next(x for x in c if x["cl"] == "8")["sal"]
    assert abs(l8 - M.dre(b, "2026-08")["acumulado"]["res_liq"]) < 0.005      # resultado liquido acumulado

def test_balanco_igual_agosto(csvd, mensal):
    """Contas de ativo/passivo do acumulado = balancete mensal de agosto (exceto 2.4.13, onde o resultado ainda esta na conta 8)."""
    _, c = csvd; b, d = mensal
    ag = {x["id"]: x for x in d["2026-08"]["contas"]}
    for x in c:
        if x["cl"][0] in "12" and not x["cl"].startswith("2.4.13") and x["cl"] not in ("2", "2.4"):
            assert abs(ag[x["id"]]["sal"] - x["sal"]) < 0.005 if x["id"] in ag else abs(x["sal"]) < 0.005, x["cl"]

def test_explica_693659_62(mensal):
    """Linha 8 do balancete mensal de agosto (mes encerrado) conta a receita financeira como despesa e as deducoes como receita:
    diferenca para a DRE = 2*(receitas financeiras + deducoes) + ajuste 7,40."""
    b, d = mensal
    ag = M.dre(b, "2026-08")["mes"]
    l8 = next(x for x in d["2026-08"]["contas"] if x["cl"] == "8")["deb"]     # mes encerrado: valor esta no debito
    dif = round(abs(l8) - abs(ag["res_liq"]), 2)
    assert dif == 693659.62
    assert abs(2 * ag["rec_fin"] + 2 * ag["deducoes"] + ag["ajustes"] - dif) < 0.005


def test_pdf_do_acumulado_tem_os_mesmos_numeros_do_csv(pasta, csvd):
    """O balancete em PDF (opcional na pasta) deve dar exatamente os mesmos numeros do CSV; so' os nomes longos vem cortados no PDF."""
    import importador_pdf as P
    f = pasta / "01 a 08.2026 - Balancete - sem assinatura Alex.pdf"
    if not f.exists():
        pytest.skip("PDF do acumulado nao esta na pasta")
    cab_csv, c_csv = csvd
    cab, c_pdf = P.ler_pdf(str(f))
    assert (cab["tipo"], cab["periodo"], cab["cnpj"]) == (cab_csv["tipo"], cab_csv["periodo"], cab_csv["cnpj"])
    assert len(c_pdf) == len(c_csv) and not cab["ignoradas"]
    for a, b in zip(c_pdf, c_csv):
        assert (a["id"], a["cl"], a["sint"]) == (b["id"], b["cl"], b["sint"])
        assert all(abs(a[k] - b[k]) < 0.005 for k in ("ant", "deb", "cred", "sal"))
        assert " ".join(b["nome"].split()).startswith(" ".join(a["nome"].split())[:30])


def test_pdf_mensal_assinado_confere(pasta):
    import importador_pdf as P
    f = pasta / "08.2026 - Balancete - [assinado].pdf"
    if not f.exists():
        pytest.skip("PDF assinado de agosto nao esta na pasta")
    cab, contas = P.ler_pdf(str(f))
    assert (cab["tipo"], cab["periodo"], len(contas)) == ("MENSAL", "2026-08", 187)
    assert all(c["ok"] for c in M.conferencias_arquivo(contas, None, "2026-08"))


def test_relatorio_pdf_com_dados_reais_tem_os_valores_do_excel(b, wb):
    """O PDF gerado a partir dos balancetes traz o resultado e o caixa do Excel da contadora (centavo a centavo)."""
    import io
    import pdfplumber
    import relatorio_dados as RD
    import relatorio_pdf as RP
    mes = "2026-08"
    bp, dr = M.balanco(b, mes), M.dre(b, mes)
    emp = {"id": 1, "codigo": "ANASTACIO", "razao_social": "Anastácio Transmissora de Energia S.A.", "cnpj": "54.800.488/0001-60"}
    ctx = RD.montar(emp, b, mes, bp, dr, b.p[mes])
    pdf = RP.gerar_pdf(ctx, {}, "RASCUNHO", "01/01/2026 00:00")
    with pdfplumber.open(io.BytesIO(pdf)) as p:
        assert len(p.pages) == 11
        todo = "\n".join(pg.extract_text() or "" for pg in p.pages)
    ws = wb["Balanço Patrimonial"]
    caixa = next(ws.cell(r, 5).value for r in range(10, 54) if ws.cell(r, 2).value and "Caixa e equivalentes" in str(ws.cell(r, 2).value))
    assert RP.fnum(caixa) in todo
    assert RP.fnum(dr["acumulado"]["res_liq"]) in todo
    # "Fornecedores — 10 maiores": 10 nomes + demais, como no Excel (o 10o e' a Omicron, R$ 442.800,00)
    assert "Omicron" in todo and RP.fnum(442800.0) in todo


def test_painel_serie_com_dados_reais(b):
    import painel
    linhas = painel.serie_mensal(b)
    assert len(linhas) == len(b.ordem) and linhas[0]["periodo"].endswith("-01")
    assert abs(sum(l["resultado_mes"] for l in linhas) - linhas[-1]["resultado_acum"]) < 0.05
    assert abs(linhas[-1]["resultado_acum"] - M.dre(b, b.ordem[-1])["acumulado"]["res_liq"]) < 0.005
    assert linhas[-1]["periodo"] == b.ordem[-1] and linhas[-1]["caixa"] > 0
