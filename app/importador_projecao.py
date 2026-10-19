# -*- coding: utf-8 -*-
"""
GDF V2 - leitor da planilha de PROJECAO DE RECEITAS DAS OBRAS (xlsx, modelo da pasta UPAR: 'Projecoes-Receitas-Obras.xlsx').
Colunas: CONTRATO, CLIENTE, OBRA, CATEGORIA, NATUREZA, EMPRESA / FORNECEDOR, <mes da projecao, ex. 'set/26'>, CUSTOS, BDI; a ultima linha e' o TOTAL.
Devolve o mes da projecao, as linhas (receita prevista, custos previstos, BDI previsto) e avisos de consistencia (BDI = receita - custos; soma = linha TOTAL).
Le so' valores (openpyxl, data_only); o arquivo nunca e' alterado. Sem streamlit, sem banco.
"""
from __future__ import annotations

import io
import re
import unicodedata

_MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}
_RE_MES = re.compile(r"^\s*([a-zç]{3})\s*[/\-]\s*(\d{2}|\d{4})\s*$", re.I)


class ErroProjecao(ValueError):
    """Mensagem em portugues simples, pronta para mostrar na tela."""


def _sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(s or "")) if not unicodedata.combining(c))


def normalizar(s: str) -> str:
    """Texto em maiusculas, sem acento, so' letras e numeros separados por espaco (para casar nomes de obra)."""
    return re.sub(r"[^A-Z0-9]+", " ", _sem_acento(s).upper()).strip()


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        raise ErroProjecao(f"Valor que nao consegui ler como numero: “{v}”.")


def ler_projecao_xlsx(conteudo: bytes) -> dict:
    """Le a planilha. Devolve {'mes': 'AAAA-MM', 'linhas': [{linha, contrato, cliente, obra, categoria, natureza, empresa, valor_mes, custos, bdi}],
    'total': {valor_mes, custos, bdi} (somado das linhas), 'total_planilha': {...}|None, 'avisos': [...]}."""
    import openpyxl
    if not conteudo or conteudo[:2] != b"PK":
        raise ErroProjecao("Este arquivo nao e' uma planilha do Excel (.xlsx). Envie a planilha 'Projeções-Receitas-Obras'.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception:
        raise ErroProjecao("Nao consegui abrir a planilha. Confira se o arquivo nao esta protegido ou corrompido.")
    ws = wb[wb.sheetnames[0]]
    linhas_xl = list(ws.iter_rows(values_only=True))
    # cabecalho: primeira linha que tem CONTRATO e CUSTOS e BDI
    ih, cab = None, None
    for i, r in enumerate(linhas_xl[:15]):
        nomes = [normalizar(c) for c in r]
        if "CONTRATO" in nomes and "CUSTOS" in nomes and "BDI" in nomes:
            ih, cab = i, nomes
            break
    if ih is None:
        raise ErroProjecao("Nao achei o cabecalho da planilha (colunas CONTRATO, ..., CUSTOS, BDI). Confira se e' a planilha de projecao de receitas.")
    col = {n: j for j, n in enumerate(cab) if n}
    # coluna do mes: a que parece 'set/26'
    jmes, mes = None, None
    for j, c in enumerate(linhas_xl[ih]):
        m = _RE_MES.match(str(c or ""))
        if m and _sem_acento(m.group(1)).lower() in _MESES:
            ano = int(m.group(2)) + (2000 if len(m.group(2)) == 2 else 0)
            jmes, mes = j, f"{ano}-{_MESES[_sem_acento(m.group(1)).lower()]:02d}"
            break
    if jmes is None:
        raise ErroProjecao("Nao achei a coluna do mes da projecao (cabecalho como 'set/26').")
    jemp = next((j for n, j in col.items() if n.startswith("EMPRESA")), None)
    g = lambda r, nome: (r[col[nome]] if nome in col and col[nome] < len(r) else None)
    linhas, tot_pl, avisos = [], None, []
    for k, r in enumerate(linhas_xl[ih + 1:], start=ih + 2):
        if all(c in (None, "") for c in r):
            continue
        rotulo = normalizar(r[jemp]) if jemp is not None and jemp < len(r) else ""
        valor = _num(r[jmes] if jmes < len(r) else None)
        if rotulo == "TOTAL" and not g(r, "CONTRATO"):
            tot_pl = {"valor_mes": round(valor, 2), "custos": round(_num(g(r, "CUSTOS")), 2), "bdi": round(_num(g(r, "BDI")), 2)}
            continue
        if not str(g(r, "OBRA") or "").strip() and not str(g(r, "CONTRATO") or "").strip():
            continue
        custos, bdi = _num(g(r, "CUSTOS")), _num(g(r, "BDI"))
        if abs(valor - custos - bdi) > 0.05:
            avisos.append(f"Linha {k}: receita ({valor:,.2f}) nao fecha com custos + BDI ({custos + bdi:,.2f}).")
        linhas.append({"linha": k, "contrato": str(g(r, "CONTRATO") or "").strip(), "cliente": str(g(r, "CLIENTE") or "").strip(),
                       "obra": str(g(r, "OBRA") or "").strip(), "categoria": str(g(r, "CATEGORIA") or "").strip(),
                       "natureza": str(g(r, "NATUREZA") or "").strip(), "empresa": str(r[jemp] or "").strip() if jemp is not None else "",
                       "valor_mes": round(valor, 2), "custos": round(custos, 2), "bdi": round(bdi, 2)})
    if not linhas:
        raise ErroProjecao("A planilha nao tem linhas de projecao abaixo do cabecalho.")
    total = {k: round(sum(l[k] for l in linhas), 2) for k in ("valor_mes", "custos", "bdi")}
    if tot_pl is not None:
        for k, nome in (("valor_mes", "receita"), ("custos", "custos"), ("bdi", "BDI")):
            if abs(total[k] - tot_pl[k]) > 0.05:
                avisos.append(f"A soma das linhas ({nome}) e' {total[k]:,.2f}, diferente da linha TOTAL da planilha ({tot_pl[k]:,.2f}).")
    return {"mes": mes, "linhas": linhas, "total": total, "total_planilha": tot_pl, "avisos": avisos}


def projecao_por_obra(linhas: list) -> list:
    """Soma a projecao por nome de obra (normalizado): [{'obra', 'chave', 'receita', 'custos', 'bdi', 'cliente'}], maior receita primeiro."""
    por = {}
    for l in linhas:
        k = normalizar(l["obra"])
        a = por.setdefault(k, {"obra": l["obra"].strip(), "chave": k, "cliente": l["cliente"], "receita": 0.0, "custos": 0.0, "bdi": 0.0})
        a["receita"] += l["valor_mes"]; a["custos"] += l["custos"]; a["bdi"] += l["bdi"]
    return sorted(({**a, "receita": round(a["receita"], 2), "custos": round(a["custos"], 2), "bdi": round(a["bdi"], 2)} for a in por.values()), key=lambda a: -a["receita"])


def casar_obra_dre(chave_projecao: str, obras_dre: dict) -> str | None:
    """Acha o codigo da obra do DRE para o nome da projecao: {'103': 'ANASTACIO - LOTE 11 - 001/2024', ...}. A projecao usa nomes curtos ('LOTE 11',
    'COAMO'): casa quando TODAS as palavras do nome da projecao aparecem no nome da obra do DRE (palavra inteira: 'OPGW II' nao casa com 'OPGW III').
    Devolve None se nenhuma ou mais de uma obra casa (a tela lista essas como 'sem correspondencia')."""
    palavras = set(chave_projecao.split())
    if not palavras:
        return None
    ach = [cod for cod, nome in obras_dre.items() if palavras <= set(normalizar(nome).split())]
    return ach[0] if len(ach) == 1 else None
