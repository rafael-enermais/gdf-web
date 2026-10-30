# -*- coding: utf-8 -*-
"""GDF V2 - Relatorio do Painel da Diretoria em WORD (.docx), v0.63.0. Desenha, em texto editavel, a MESMA estrutura (`relatorio_diretoria_dados.montar`) que o PDF usa:
capa com periodo e filtros, uma secao por aba, uma subsecao por informacao, cartoes (kpis), tabelas, notas e a observacao da aba.
Graficos de barras viram uma tabela de valores (rotulo x valor), que o Word edita e transforma em grafico se quiser. Nada vai ao banco."""
from __future__ import annotations

import io

from relatorio_fmt import fnum, mmn

_FMT = {
    "mi": lambda v: mmn(v, 1),
    "mi2": lambda v: mmn(v, 2),
    "int": lambda v: fnum(v, 0),
    "p1": lambda v: fnum(v, 1),
    "mil": lambda v: fnum(v / 1000.0, 0),
}
_AZUL = "0F2A4A"
_CINZA = "F2F4F7"
_POR_LINHA_BARRAS = 12


def _cel_txt(v) -> str:
    return str(v[0] if isinstance(v, tuple) else v)


def _sombra(cell, cor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), cor)
    tc.append(sh)


def _escreve(cell, texto: str, negrito: bool = False, dir_: bool = False, cor=None, tam: float = 8.0, italico: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if dir_:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(str(texto))
    r.bold = negrito
    r.italic = italico
    r.font.size = Pt(tam)
    if cor:
        r.font.color.rgb = RGBColor.from_string(cor)


def _tabela_simples(doc, cabecalho: list | None, linhas: list, alin: list | None = None, negrito_linhas: set | None = None, tam: float = 8.0):
    n = len(cabecalho) if cabecalho else (len(linhas[0]) if linhas else 1)
    t = doc.add_table(rows=0, cols=n)
    t.style = "Table Grid"
    alin = alin or (["l"] + ["r"] * (n - 1))
    if cabecalho:
        cs = t.add_row().cells
        for i, h in enumerate(cabecalho):
            _escreve(cs[i], _cel_txt(h), negrito=True, dir_=(alin[i] == "r"), cor="FFFFFF", tam=tam)
            _sombra(cs[i], _AZUL)
    for k, ln in enumerate(linhas):
        cs = t.add_row().cells
        for i in range(n):
            v = ln[i] if i < len(ln) else ""
            _escreve(cs[i], _cel_txt(v), negrito=bool(negrito_linhas and k in negrito_linhas), dir_=(alin[i] == "r"), tam=tam)
            if negrito_linhas and k in negrito_linhas:
                _sombra(cs[i], _CINZA)
    return t


def _kpis(doc, e: dict) -> None:
    itens = e.get("itens") or []
    por = max(1, int(e.get("n", 4)))
    for i in range(0, len(itens), por):
        grupo = itens[i:i + por]
        t = doc.add_table(rows=1, cols=por)
        t.style = "Table Grid"
        for j, it in enumerate(grupo):
            c = t.rows[0].cells[j]
            _escreve(c, str(it.get("rot", "")), tam=7.5, cor="5B6770")
            p = c.add_paragraph()
            r = p.add_run(str(it.get("val", "")))
            r.bold = True
            from docx.shared import Pt
            r.font.size = Pt(13)
            if it.get("sub"):
                p2 = c.add_paragraph()
                r2 = p2.add_run(str(it["sub"]))
                r2.font.size = Pt(7)
        doc.add_paragraph()


def _barras(doc, e: dict) -> None:
    labels, vals = list(e.get("labels") or []), list(e.get("vals") or [])
    if not vals:
        return
    fmt = _FMT.get(e.get("fmt") or "mi", _FMT["mi"])
    if e.get("titulo"):
        p = doc.add_paragraph()
        r = p.add_run(f"{e['titulo']} ({e.get('unidade') or 'R$ MM'})")
        r.bold = True
    for i in range(0, len(vals), _POR_LINHA_BARRAS):
        lb = labels[i:i + _POR_LINHA_BARRAS]
        vs = vals[i:i + _POR_LINHA_BARRAS]
        _tabela_simples(doc, [str(x) for x in lb], [[("–" if v is None else fmt(v)) for v in vs]], ["r"] * len(lb), tam=7.5)
    doc.add_paragraph()


def _tab(doc, e: dict) -> None:
    if e.get("titulo"):
        p = doc.add_paragraph()
        p.add_run(str(e["titulo"])).bold = True
    rows = e.get("rows") or []
    linhas = [cel for _tipo, cel in rows]
    negr = {k for k, (tipo, _c) in enumerate(rows) if tipo in ("sub", "tot", "sec")}
    _tabela_simples(doc, e.get("head"), linhas, e.get("alin"), negr)
    doc.add_paragraph()


def gerar_docx(dados: dict) -> bytes:
    """Word (bytes) do relatorio. `dados` = relatorio_diretoria_dados.montar(...)."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt, RGBColor
    meta = dados["meta"]
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)
    doc.core_properties.title = f"Relatório do Painel da Diretoria — {meta.get('per_rot', '')}"
    doc.core_properties.author = "Grupo EnerMais — GDF"

    h = doc.add_heading("Relatório do Painel da Diretoria", level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor.from_string(_AZUL)
    doc.add_paragraph(f"Período: {meta.get('per_rot', '')}")
    for ln in meta.get("filtros", []) or []:
        doc.add_paragraph(str(ln))
    doc.add_paragraph(f"Gerado em {meta.get('gerado_em', '')}" + (f" por {meta['usuario']}" if meta.get("usuario") else "")
                      + ". Informações do Sienge (API) e dos arquivos importados no painel. Valores em R$ salvo indicação.")

    for ab in dados.get("abas", []):
        doc.add_page_break()
        doc.add_heading(str(ab.get("rotulo") or ab.get("aba")), level=1)
        for bl in ab.get("blocos", []):
            doc.add_heading(str(bl.get("titulo") or ""), level=2)
            for e in bl.get("els", []):
                t = e.get("t")
                if t == "kpis":
                    _kpis(doc, e)
                elif t == "tab":
                    _tab(doc, e)
                elif t == "barras":
                    _barras(doc, e)
                elif t == "sec":
                    doc.add_heading(str(e.get("texto") or ""), level=3)
                elif t == "nota":
                    p = doc.add_paragraph()
                    r = p.add_run(str(e.get("texto") or ""))
                    r.italic = True
                    r.font.size = Pt(8)
        if ab.get("obs"):
            t = doc.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            _escreve(t.rows[0].cells[0], "Observação: " + ab["obs"], tam=8.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
