# -*- coding: utf-8 -*-
"""GDF V2 - cargas do painel "Folha PJ": (1) mapa de colaboradores PJ (a planilha da folha, ou o arquivo "Novos colaboradores PJ"), (2) Flash (um xlsx por empresa)
e (3) cadastro dos credores da diretoria. So' le planilhas (openpyxl) e grava nas tabelas pj_* (SQL-40). O CPF NUNCA e' gravado: so' o hash (HMAC-SHA256 com segredo do app).

Em linguagem simples: o app acha sozinho a linha do cabecalho pelo NOME das colunas (nao importa a ordem nem a linha), le so' as pessoas PJ, e mostra o que entendeu antes de salvar.
Planilha protegida por senha: o app nao abre; salve uma copia sem senha (Arquivo > Informacoes > Proteger pasta de trabalho > Criptografar com senha > apagar a senha)."""
from __future__ import annotations

import datetime as dt
import io
import re

import openpyxl
from psycopg2.extras import execute_values

import db
import folha_pj_dados as FP

ASSINATURA_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"          # xlsx protegido por senha vira um arquivo OLE (nao e' um zip)


class ErroFolhaPJ(Exception):
    """Erro de leitura explicado em linguagem simples."""


def _abrir(conteudo: bytes):
    if conteudo[:8] == ASSINATURA_OLE:
        raise ErroFolhaPJ("A planilha está protegida por senha e o app não consegue abrir. Salve uma cópia sem senha (Arquivo > Informações > Proteger Pasta de Trabalho > "
                          "Criptografar com Senha > apagar a senha) e envie essa cópia.")
    try:
        return openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception as exc:                                         # noqa: BLE001
        raise ErroFolhaPJ(f"Não consegui abrir o arquivo como planilha xlsx ({type(exc).__name__}).") from exc


def abas(conteudo: bytes) -> list:
    wb = _abrir(conteudo)
    return [ws.title for ws in wb.worksheets]


def aba_padrao(nomes: list) -> str:
    """A aba do mes mais recente ('MM-AAAA') se houver; senao a primeira."""
    meses = sorted((n for n in nomes if re.fullmatch(r"\d{2}-\d{4}", n)), key=lambda n: (n[3:], n[:2]))
    return meses[-1] if meses else nomes[0]


def _n(s) -> str:
    return FP.sem_acento(re.sub(r"\s+", " ", str(s or "").replace("\n", " "))).strip()


def _achar_cabecalho(linhas: list, exige: tuple) -> tuple:
    """(indice da linha, {nome normalizado: coluna}) da 1a linha dos 40 primeiros que tem todas as colunas de `exige` (por inicio de nome)."""
    for i, r in enumerate(linhas[:40]):
        nomes = [_n(c) for c in r]
        if all(any(n.startswith(e) for n in nomes) for e in exige):
            return i, nomes
    return -1, []


def _col(nomes: list, *padroes, exato=False):
    for p in padroes:
        for j, n in enumerate(nomes):
            if (n == p) if exato else (p in n):
                return j
    return None


def _dig(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


def _int(v):
    try:
        return int(float(str(v).replace(",", ".").strip()))
    except (TypeError, ValueError):
        return None


def _data(v):
    if isinstance(v, dt.datetime):
        return v.date()
    return v if isinstance(v, dt.date) else None


_PALAVRAS_EMPRESA = (("smg", 2), ("renov", 4), ("construtora", 5), ("engenharia", 3), ("solu", 6), ("energia", 1))        # palavra no nome do arquivo -> codigo da empresa (FP.EMPRESAS)


def empresa_do_nome(nome_arquivo: str) -> str | None:
    """'FLASH  RENOVAVEIS.xlsx' -> 'ENERMAIS RENOVAVEIS LTDA'. None se o nome do arquivo nao diz a empresa (ai o usuario escolhe)."""
    n = FP.sem_acento(nome_arquivo)
    for palavra, cod in _PALAVRAS_EMPRESA:
        if palavra in n:
            return FP.EMPRESAS[cod]
    return None


def detectar(conteudo: bytes, nome_arquivo: str = "") -> dict:
    """Descobre sozinho o que e' o arquivo, so' olhando o cabecalho: 'flash' (coluna NOME COMPLETO + premiacao/refeicao), 'mapa' (colunas NOME e COD. CREDOR),
    'novos' (o mesmo, mas e' o arquivo 'Novos colaboradores PJ') ou 'desconhecido'. Para o Flash acha a empresa pelo nome do arquivo.
    Devolve {'tipo','empresa','abas','aba'} (aba = aba do mes mais recente, so' no mapa)."""
    wb = _abrir(conteudo)
    nomes_abas = [ws.title for ws in wb.worksheets]
    def topo(ws):
        return [r for _, r in zip(range(40), ws.iter_rows(values_only=True))]
    primeira = topo(wb.worksheets[0])
    ih, h = _achar_cabecalho(primeira, ("nome completo",))
    if ih >= 0 and any(re.search(r"premia|refei|aliment", n) for n in h):
        return {"tipo": "flash", "empresa": empresa_do_nome(nome_arquivo), "abas": nomes_abas, "aba": nomes_abas[0]}
    for ws in wb.worksheets:
        lin = topo(ws)
        i2, h2 = _achar_cabecalho(lin, ("nome",))
        if i2 >= 0 and _col(h2, "cod credor", "cod. credor") is not None:
            novos = "novos" in FP.sem_acento(nome_arquivo) or "novos colaboradores" in FP.sem_acento(ws.title)
            return {"tipo": "novos" if novos else "mapa", "empresa": None, "abas": nomes_abas, "aba": ws.title if novos else aba_padrao(nomes_abas)}
    return {"tipo": "desconhecido", "empresa": None, "abas": nomes_abas, "aba": nomes_abas[0]}


def resumo_salvo(conn) -> dict:
    """O que ja' esta gravado (para a tela mostrar 'salvo' de verdade): mapa (quantos e quando) e Flash por mes/empresa (so' as cargas ativas)."""
    with conn.cursor() as cur:
        cur.execute("SELECT count(*), max(alterado_em) FROM pj_colaborador WHERE ativo")
        n, quando = cur.fetchone()
        cur.execute("""SELECT to_char(mes,'YYYY-MM'), empresa, count(*), COALESCE(sum(premiacao),0), COALESCE(sum(alimentacao),0), max(importado_por), max(importado_em)
                       FROM pj_flash WHERE ativo GROUP BY mes, empresa ORDER BY mes DESC, empresa""")
        fl = [{"mes": m, "empresa": e, "pessoas": p, "premiacao": float(a), "alimentacao": float(b), "quem": q, "quando": w} for m, e, p, a, b, q, w in cur.fetchall()]
    return {"colaboradores": int(n or 0), "quando": quando, "flash": fl}


def ler_mapa_colaboradores(conteudo: bytes, segredo: str, aba: str | None = None) -> dict:
    """Le a planilha da folha (aba do mes) ou o arquivo 'Novos colaboradores PJ'. Devolve {'aba','linhas':[{credor_id,nome,funcao,centro,empresa,cpf_hash,admissao}],'avisos':[...]}.
    So' entram linhas de vinculo PJ (ou sem coluna de vinculo) com codigo de credor do Sienge."""
    wb = _abrir(conteudo)
    nomes_abas = [ws.title for ws in wb.worksheets]
    aba = aba if aba in nomes_abas else aba_padrao(nomes_abas)
    linhas = list(wb[aba].iter_rows(values_only=True))
    ih, h = _achar_cabecalho(linhas, ("nome",))
    if ih < 0 or _col(h, "cod credor", "cod. credor", "cod credor sieng") is None:
        raise ErroFolhaPJ(f"Não achei, na aba “{aba}”, o cabeçalho com as colunas NOME e CÓD. CREDOR SIENGE.")
    c = {"nome": _col(h, "nome", exato=True), "cod": _col(h, "cod credor", "cod. credor"), "credor": _col(h, "credor sienge", exato=True), "cpf": _col(h, "cpf", exato=True),
         "funcao": _col(h, "funcao"), "centro": _col(h, "centro de custo", "cento de custo"), "empresa": _col(h, "empresa", exato=True), "vinculo": _col(h, "vinculo"),
         "adm": _col(h, "admissao")}
    if c["nome"] is None:
        raise ErroFolhaPJ("Não achei a coluna NOME.")
    out, avisos, sem_cod, nao_pj = [], [], 0, 0
    for r in linhas[ih + 1:]:
        g = lambda k: (r[c[k]] if c[k] is not None and c[k] < len(r) else None)        # noqa: E731
        nome = str(g("nome") or "").strip()
        if not nome:
            continue
        if c["vinculo"] is not None and not _n(g("vinculo")).startswith("pj"):
            nao_pj += 1
            continue
        cod = _int(g("cod"))
        if cod is None:
            sem_cod += 1
            continue
        out.append({"credor_id": cod, "nome": nome, "funcao": str(g("funcao") or "").strip() or None, "centro": str(g("centro") or "").strip() or None,
                    "empresa": str(g("empresa") or "").strip() or None, "cpf_hash": FP.hash_cpf(g("cpf"), segredo), "admissao": _data(g("adm"))})
    unicos, repetidos = {}, 0
    for x in out:                                                    # a mesma pessoa em duas linhas (dois centros): fica a 1a e completa o que faltar
        o = unicos.get(x["credor_id"])
        if o is None:
            unicos[x["credor_id"]] = x
        else:
            repetidos += 1
            for k2 in ("funcao", "centro", "empresa", "cpf_hash", "admissao"):
                o[k2] = o[k2] or x[k2]
    out = list(unicos.values())
    if repetidos:
        avisos.append(f"{repetidos} linhas repetiam um credor já lido (a mesma pessoa em mais de uma linha): foi mantida a primeira.")
    if nao_pj:
        avisos.append(f"{nao_pj} linhas que não são PJ (CLT, estagiário etc.) foram ignoradas.")
    if sem_cod:
        avisos.append(f"{sem_cod} linhas PJ sem código de credor foram ignoradas (a diretoria agrupada não tem código).")
    sem_cpf = sum(1 for x in out if not x["cpf_hash"])
    if sem_cpf:
        avisos.append(f"{sem_cpf} colaboradores sem CPF válido: o Flash deles será cruzado só pelo nome.")
    if not out:
        raise ErroFolhaPJ(f"Nenhuma pessoa PJ com código de credor na aba “{aba}”.")
    return {"aba": aba, "linhas": out, "avisos": avisos}


def ler_flash(conteudo: bytes, segredo: str) -> dict:
    """Le um xlsx do Flash (um por empresa). Colunas achadas pelo nome: NOME COMPLETO, CPF, PREMIACAO NO CARTAO, e REFEICAO / ALIMENTACAO (soma das colunas que existirem).
    Devolve {'linhas':[{nome, nome_chave, cpf_hash, premiacao, alimentacao}], 'total_premiacao', 'total_alimentacao', 'avisos'}."""
    wb = _abrir(conteudo)
    linhas = list(wb.worksheets[0].iter_rows(values_only=True))
    ih, h = _achar_cabecalho(linhas, ("nome completo",))
    if ih < 0:
        raise ErroFolhaPJ("Não achei a coluna NOME COMPLETO no arquivo do Flash.")
    i_nome, i_cpf = _col(h, "nome completo"), _col(h, "cpf", exato=True)
    i_prem = [j for j, n in enumerate(h) if "premia" in n]
    i_ali = [j for j, n in enumerate(h) if re.search(r"refei|aliment", n)]
    if not i_prem and not i_ali:
        raise ErroFolhaPJ("Não achei as colunas de premiação no cartão e de refeição/alimentação.")
    num = lambda v: float(str(v).replace(",", ".")) if str(v or "").strip().replace(",", ".").replace(".", "", 1).replace("-", "", 1).isdigit() else 0.0     # noqa: E731
    out = []
    for r in linhas[ih + 1:]:
        nome = str(r[i_nome] or "").strip() if i_nome is not None and i_nome < len(r) else ""
        if not nome:
            continue
        p = sum(num(r[j]) for j in i_prem if j < len(r))
        a = sum(num(r[j]) for j in i_ali if j < len(r))
        out.append({"nome": nome, "nome_chave": FP.chave_nome(nome), "cpf_hash": FP.hash_cpf(r[i_cpf] if i_cpf is not None and i_cpf < len(r) else None, segredo),
                    "premiacao": round(p, 2), "alimentacao": round(a, 2)})
    avisos = []
    extras = [h[j] for j, n in enumerate(h) if re.search(r"flexivel|mobilidade", n) and any(num(r[j]) for r in linhas[ih + 1:] if j < len(r))]
    if extras:
        avisos.append("O arquivo tem valores em colunas que o painel não usa (" + ", ".join(sorted({e.upper() for e in extras})) + ").")
    return {"linhas": out, "total_premiacao": round(sum(x["premiacao"] for x in out), 2), "total_alimentacao": round(sum(x["alimentacao"] for x in out), 2), "avisos": avisos}


def cruzar_flash(flash: list, colaboradores: dict, fantasias: dict | None = None) -> dict:
    """Casa cada pessoa do Flash com um colaborador PJ: primeiro pelo hash do CPF, depois pelo nome (cadastro do mapa e nome fantasia do Sienge).
    colaboradores = {credor_id: {'cpf_hash','chave'}}; fantasias = {credor_id: nome fantasia}. Nome repetido em mais de um credor nao casa (ambiguo).
    Devolve {'casadas':[{credor_id, premiacao, alimentacao, via}], 'sem_pj': n, 'ambiguas': n}."""
    por_cpf = {}
    for cid, i in colaboradores.items():
        if i.get("cpf_hash"):
            por_cpf.setdefault(i["cpf_hash"], []).append(cid)
    por_nome = {}
    for cid, i in colaboradores.items():
        for n in (i.get("chave"), FP.chave_nome((fantasias or {}).get(cid))):
            if n:
                por_nome.setdefault(n, set()).add(cid)
    for cid, n in (fantasias or {}).items():
        if FP.chave_nome(n):
            por_nome.setdefault(FP.chave_nome(n), set()).add(cid)
    casadas, sem, amb = [], 0, 0
    for f in flash:
        cands, via = [], ""
        if f["cpf_hash"] and f["cpf_hash"] in por_cpf:
            cands, via = por_cpf[f["cpf_hash"]], "CPF"
        elif f["nome_chave"] in por_nome:
            cands, via = sorted(por_nome[f["nome_chave"]]), "nome"
        if not cands:
            sem += 1
        elif len(cands) > 1:
            amb += 1
        elif f["premiacao"] or f["alimentacao"]:
            casadas.append({"credor_id": cands[0], "premiacao": f["premiacao"], "alimentacao": f["alimentacao"], "via": via})
    return {"casadas": casadas, "sem_pj": sem, "ambiguas": amb}


# ------------------------------------------------------------------ gravacao (tabelas do SQL-40)
def gravar_colaboradores(conn, linhas: list, mes: str, usuario: str | None, novos: bool = False) -> dict:
    """Upsert do mapa. 'novos' = o arquivo e' o 'Novos colaboradores PJ': marca novo_desde = 1o dia do mes (so' se a pessoa ainda nao tinha data)."""
    nd = dt.date.fromisoformat(f"{mes}-01") if novos else None
    rows = [(l["credor_id"], l["nome"], FP.chave_nome(l["nome"]), l["funcao"], l["centro"], l["empresa"], l["cpf_hash"], nd, usuario) for l in linhas]
    with db.transacao(conn), conn.cursor() as cur:
        cur.execute("SELECT credor_id FROM pj_colaborador")
        antes = {r[0] for r in cur.fetchall()}
        execute_values(cur, """INSERT INTO pj_colaborador (credor_id, nome_exibicao, nome_chave, funcao, centro_origem, empresa_origem, cpf_hash, novo_desde, alterado_por) VALUES %s
                               ON CONFLICT (credor_id) DO UPDATE SET nome_exibicao=EXCLUDED.nome_exibicao, nome_chave=EXCLUDED.nome_chave,
                                   funcao=COALESCE(EXCLUDED.funcao, pj_colaborador.funcao), centro_origem=COALESCE(EXCLUDED.centro_origem, pj_colaborador.centro_origem),
                                   empresa_origem=COALESCE(EXCLUDED.empresa_origem, pj_colaborador.empresa_origem), cpf_hash=COALESCE(EXCLUDED.cpf_hash, pj_colaborador.cpf_hash),
                                   novo_desde=COALESCE(pj_colaborador.novo_desde, EXCLUDED.novo_desde), ativo=true, alterado_por=EXCLUDED.alterado_por, alterado_em=now()""", rows)
    ids = {l["credor_id"] for l in linhas}
    return {"novos": len(ids - antes), "atualizados": len(ids & antes)}


def colaboradores_cadastrados(conn) -> dict:
    with conn.cursor() as cur:
        cur.execute("SELECT credor_id, cpf_hash, nome_chave FROM pj_colaborador WHERE ativo")
        return {int(i): {"cpf_hash": h, "chave": k} for i, h, k in cur.fetchall()}


def fantasias_dos(conn, ids: list) -> dict:
    if not ids:
        return {}
    with conn.cursor() as cur:
        cur.execute("SELECT id, nome_fantasia FROM pnl_credor_sienge WHERE id = ANY(%s)", (list(ids),))
        return {int(i): (n or "").strip() for i, n in cur.fetchall()}


def gravar_flash(conn, mes: str, empresa: str, casadas: list, usuario: str | None) -> dict:
    """Troca o Flash do mes para a empresa: desativa o que havia e grava o novo (nada e apagado)."""
    ini = f"{mes}-01"
    por_cred = {}
    for c in casadas:                                             # a mesma pessoa em duas linhas do arquivo soma
        a = por_cred.setdefault(c["credor_id"], [0.0, 0.0])
        a[0] += c["premiacao"]
        a[1] += c["alimentacao"]
    with db.transacao(conn), conn.cursor() as cur:
        cur.execute("UPDATE pj_flash SET ativo=false WHERE mes=%s::date AND empresa=%s AND ativo", (ini, empresa))
        if por_cred:
            execute_values(cur, """INSERT INTO pj_flash (mes, empresa, credor_id, premiacao, alimentacao, ativo, importado_por) VALUES %s
                                   ON CONFLICT (mes, empresa, credor_id) DO UPDATE SET premiacao=EXCLUDED.premiacao, alimentacao=EXCLUDED.alimentacao, ativo=true,
                                       importado_por=EXCLUDED.importado_por, importado_em=now()""",
                           [(ini, empresa, cid, round(v[0], 2), round(v[1], 2), True, usuario) for cid, v in por_cred.items()])
    return {"pessoas": len(por_cred), "premiacao": round(sum(v[0] for v in por_cred.values()), 2), "alimentacao": round(sum(v[1] for v in por_cred.values()), 2)}


def diretoria_listar(conn) -> list:
    with conn.cursor() as cur:
        cur.execute("SELECT credor_id, COALESCE(nome_exibicao,''), alterado_por, alterado_em FROM pj_diretoria_credor WHERE ativo ORDER BY credor_id")
        return [{"credor_id": int(i), "nome": n, "quem": q, "quando": w} for i, n, q, w in cur.fetchall()]


def diretoria_adicionar(conn, credor_id: int, nome: str | None, usuario: str | None) -> None:
    with db.transacao(conn), conn.cursor() as cur:
        cur.execute("""INSERT INTO pj_diretoria_credor (credor_id, nome_exibicao, ativo, alterado_por) VALUES (%s,%s,true,%s)
                       ON CONFLICT (credor_id) DO UPDATE SET nome_exibicao=COALESCE(EXCLUDED.nome_exibicao, pj_diretoria_credor.nome_exibicao), ativo=true,
                           alterado_por=EXCLUDED.alterado_por, alterado_em=now()""", (int(credor_id), (nome or None), usuario))


def diretoria_remover(conn, credor_id: int, usuario: str | None) -> None:
    with db.transacao(conn), conn.cursor() as cur:
        cur.execute("UPDATE pj_diretoria_credor SET ativo=false, alterado_por=%s, alterado_em=now() WHERE credor_id=%s", (usuario, int(credor_id)))
